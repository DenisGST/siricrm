"""Движок заполнения .docx-шаблонов плейсхолдерами вида {key}.

Особенность Word: один логический текст («{дата рождения}») часто разбит на
несколько run'ов (<w:r>) из-за правок/проверки орфографии. Поэтому простой
поиск-замена по run'у не сработает. Решение: склеиваем текст всех run'ов
абзаца, ищем плейсхолдеры в склейке, а подставляем значение ТОЧЕЧНО — в те
run'ы, на которые пришлось совпадение, не трогая остальные.

🛑 Раньше результат писался в первый run абзаца, а остальные очищались, — из-за
этого терялось всё форматирование, кроме формата первого run'а. На практике
первым run'ом часто оказывается отступ из пробелов без оформления, и жирный
шрифт ФИО в шапке запроса пропадал (29.08.2026). Теперь формат каждого run'а
сохраняется: значение наследует оформление того run'а, где НАЧАЛСЯ плейсхолдер.
"""
import io
import re

# python-docx импортируется лениво (внутри функций) — чтобы модуль (и весь
# apps.afd, который тянется из urls) импортировался даже на образе без
# установленного python-docx (актуально между rebuild'ами на prod).

_PLACEHOLDER_RE = re.compile(r"\{[^{}\n]+\}")


def _iter_paragraphs(container):
    """Рекурсивно обходит все абзацы документа, включая абзацы в таблицах."""
    for para in getattr(container, "paragraphs", []):
        yield para
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _run_spans(runs):
    """[(начало, конец)] каждого run'а в склеенном тексте абзаца."""
    spans, pos = [], 0
    for r in runs:
        spans.append((pos, pos + len(r.text)))
        pos += len(r.text)
    return spans


def _apply_one(runs, start, end, value):
    """Заменить срез [start, end) склеенного текста на `value`.

    Значение целиком кладём в run, где совпадение началось (он и задаёт
    оформление), у остальных задетых run'ов вырезаем попавшие символы.
    """
    for idx, (s, e) in enumerate(_run_spans(runs)):
        if e <= start or s >= end:
            continue
        text = runs[idx].text
        head = text[:max(0, start - s)]
        tail = text[max(0, end - s):] if e > end else ""
        ins = value if s <= start < e else ""
        runs[idx].text = head + ins + tail


def _replace_in_paragraph(paragraph, context):
    runs = paragraph.runs
    if not runs:
        return
    if "{" not in "".join(r.text for r in runs):
        return

    # По одному плейсхолдеру за проход: после каждой замены смещения меняются,
    # поэтому склеенный текст пересобираем заново. Итераций не больше, чем
    # плейсхолдеров в абзаце.
    guard = 0
    while guard < 200:
        guard += 1
        full = "".join(r.text for r in runs)
        for m in _PLACEHOLDER_RE.finditer(full):
            key = m.group(0)[1:-1]  # без фигурных скобок
            if key not in context:
                continue  # неизвестный плейсхолдер оставляем как есть
            val = context[key]
            _apply_one(runs, m.start(), m.end(), "" if val is None else str(val))
            break
        else:
            return


def render_docx(template_bytes: bytes, context: dict) -> bytes:
    """Возвращает bytes .docx с подставленными значениями.

    context: {"placeholder_key": "value", ...} — ключи БЕЗ фигурных скобок.
    Значение None трактуется как пустая строка.
    """
    from docx import Document
    doc = Document(io.BytesIO(template_bytes))
    for para in _iter_paragraphs(doc):
        _replace_in_paragraph(para, context)
    # Заголовки/колонтитулы тоже могут содержать плейсхолдеры.
    for section in doc.sections:
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer):
            for para in getattr(hf, "paragraphs", []):
                _replace_in_paragraph(para, context)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def list_placeholders(template_bytes: bytes) -> list[str]:
    """Возвращает уникальные плейсхолдеры (без скобок) из шаблона — для UI."""
    from docx import Document
    doc = Document(io.BytesIO(template_bytes))
    found = []
    seen = set()
    for para in _iter_paragraphs(doc):
        for m in _PLACEHOLDER_RE.finditer(para.text):
            key = m.group(0)[1:-1]
            if key not in seen:
                seen.add(key)
                found.append(key)
    return found
